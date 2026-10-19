"""
Universal DOCX Document Reconstructor for Compixor Image-to-Word Universal Engine.
Conforms to PROJECT.md § Interface Contracts § 5 and Requirements R3, R4, R5, R6.

Responsibilities:
1. Spatial Sorting & Interleaving:
   Merges FormattedLine, ExtractedTable, ExtractedVisualElement, and SingleCellContainer
   into a unified reading flow ordered strictly by vertical coordinate (y-axis).
2. Word Table Grid Construction:
   - Builds genuine Microsoft Word tables with table.style = 'Table Grid'
   - Sets proportional column widths matching image ratios
   - Applies header row shading with XML fill #EBF8FF
   - Formats cell text, alignment (left/center/right), padding, and borders
   - Merges cells where col_span > 1 or row_span > 1
3. Spatial Visual Asset Embedding:
   - Embeds RGB visual crops (photos, crest logos, barcodes, signatures, stamps)
   - Converts crops to in-memory PNG byte streams
   - Preserves original aspect ratios and positions elements horizontally (left/center/right)
4. BiDi & Urdu Nastaliq Typography in Word:
   - Integrates OpenXML BiDi tags (<w:bidi w:val="1"/>, <w:rtl w:val="1"/>)
   - Binds complex script font "Jameel Noori Nastaleeq" with fallback "Noto Nastaliq Urdu"
   - Wraps boxed Urdu instructions inside a 1x1 single-cell bordered Word container table
5. In-Memory Output:
   - Direct serialization to io.BytesIO, returning raw .docx bytes with zero temp disk footprint.
"""

from __future__ import annotations

import io
import logging
from typing import List, Tuple, Optional, Dict, Any
from PIL import Image

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from pipeline.preprocessor import PreprocessedPage
from pipeline.visual_extractor import ExtractedVisualElement
from pipeline.table_extractor import ExtractedTable, TableCell
from pipeline.typography_engine import FormattedLine
from pipeline.script_router import (
    apply_bidi_paragraph_formatting,
    apply_bidi_run_formatting,
    set_cell_box_borders,
    create_urdu_container_table,
    ScriptClassifier,
    SingleCellContainer,
)

logger = logging.getLogger("compixor.pipeline.docx_reconstructor")


# ==============================================================================
# OpenXML XML Formatting Helpers
# ==============================================================================

def set_cell_shading(cell, fill_hex: str = "EBF8FF") -> None:
    """Sets background shading fill color on a Word table cell."""
    fill_hex = fill_hex.lstrip("#")
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)


def set_table_borders_and_margins(table, border_color: str = "94A3B8") -> None:
    """
    Applies explicit borders and professional cell padding to a Table Grid table.
    Ensures crisp borders across Word and LibreOffice headless rendering.
    """
    tblPr = table._tbl.tblPr

    # 1. Borders
    tblBorders = tblPr.find(qn("w:tblBorders"))
    if tblBorders is None:
        tblBorders = OxmlElement("w:tblBorders")
        tblPr.append(tblBorders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = tblBorders.find(qn(f"w:{edge}"))
        if b is None:
            b = OxmlElement(f"w:{edge}")
            tblBorders.append(b)
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "4")       # 0.5 pt
        b.set(qn("w:space"), "0")
        b.set(qn("w:color"), border_color)

    # 2. Margins / Padding (in dxa: 20 dxa = 1 pt)
    tblCellMar = tblPr.find(qn("w:tblCellMar"))
    if tblCellMar is None:
        tblCellMar = OxmlElement("w:tblCellMar")
        tblPr.append(tblCellMar)
    for edge, w_val in (("top", "120"), ("bottom", "120"), ("left", "140"), ("right", "140")):
        m = tblCellMar.find(qn(f"w:{edge}"))
        if m is None:
            m = OxmlElement(f"w:{edge}")
            tblCellMar.append(m)
        m.set(qn("w:w"), w_val)
        m.set(qn("w:type"), "dxa")


def set_table_no_borders(table) -> None:
    """Removes all borders from a layout table."""
    tblPr = table._tbl.tblPr
    tblBorders = tblPr.find(qn("w:tblBorders"))
    if tblBorders is None:
        tblBorders = OxmlElement("w:tblBorders")
        tblPr.append(tblBorders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = tblBorders.find(qn(f"w:{edge}"))
        if b is None:
            b = OxmlElement(f"w:{edge}")
            tblBorders.append(b)
        b.set(qn("w:val"), "none")


def pil_image_to_stream(img: Image.Image, format: str = "PNG") -> io.BytesIO:
    """Converts a PIL Image crop into an in-memory byte stream for add_picture."""
    stream = io.BytesIO()
    if img.mode not in ("RGB", "RGBA"):
        img = img.convert("RGB")
    img.save(stream, format=format)
    stream.seek(0)
    return stream


# ==============================================================================
# DocxReconstructor Implementation (Interface Contract 5)
# ==============================================================================

class DocxReconstructor:
    """
    Universal Document Reconstructor assembling preprocessed pages, visual assets,
    morphological tables, and typography-annotated text lines into valid .docx bytes.
    Conforms to PROJECT.md § Interface Contracts § 5.
    """

    def __init__(
        self,
        usable_width_in: float = 6.5,
        default_font: str = "Calibri",
        urdu_font_cs: str = "Jameel Noori Nastaleeq",
        urdu_font_ascii: str = "Noto Nastaliq Urdu",
        header_shading: str = "EBF8FF",
        border_color: str = "94A3B8",
    ):
        self.usable_width_in = usable_width_in
        self.default_font = default_font
        self.urdu_font_cs = urdu_font_cs
        self.urdu_font_ascii = urdu_font_ascii
        self.header_shading = header_shading
        self.border_color = border_color

    def build_document(
        self,
        pages: List[PreprocessedPage],
        visuals: List[List[ExtractedVisualElement]],
        tables: List[List[ExtractedTable]],
        text_lines: List[List[FormattedLine]],
    ) -> bytes:
        """
        Constructs valid Microsoft Word (.docx) document bytes integrating all extracted elements.
        Outputs entirely in-memory with zero disk temporary files.
        """
        doc = Document()

        # Document section page setup
        for section in doc.sections:
            section.top_margin = Inches(0.5)
            section.bottom_margin = Inches(0.5)
            section.left_margin = Inches(0.75)
            section.right_margin = Inches(0.75)

        # Document core properties
        doc.core_properties.author = "Compixor AI"
        doc.core_properties.title = "Document"
        doc.core_properties.comments = "Universal Document Extraction by Compixor AI"

        num_pages = len(pages)
        if num_pages == 0:
            buf = io.BytesIO()
            doc.save(buf)
            return buf.getvalue()

        for p_idx in range(num_pages):
            if p_idx > 0:
                doc.add_page_break()

            page = pages[p_idx]
            page_visuals = visuals[p_idx] if p_idx < len(visuals) else []
            page_tables = tables[p_idx] if p_idx < len(tables) else []
            page_lines = text_lines[p_idx] if p_idx < len(text_lines) else []

            self._build_page(doc, page, page_visuals, page_tables, page_lines)

        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)
        return buf.getvalue()

    def _build_page(
        self,
        doc: Document,
        page: PreprocessedPage,
        visuals: List[ExtractedVisualElement],
        tables: List[ExtractedTable],
        lines: List[FormattedLine],
    ) -> None:
        """Spatially sorts and interleaves elements on a single page into reading order."""
        page_w = page.width

        # 1. Group contiguous Urdu container lines into single-cell callout boxes
        layout_items: List[Dict[str, Any]] = []

        urdu_container_group: List[FormattedLine] = []
        for line in lines:
            if line.urdu_container:
                urdu_container_group.append(line)
            else:
                if urdu_container_group:
                    min_y = min(l.bbox[1] for l in urdu_container_group)
                    layout_items.append({
                        "type": "urdu_container",
                        "y": min_y,
                        "lines": list(urdu_container_group),
                    })
                    urdu_container_group.clear()
                layout_items.append({
                    "type": "line",
                    "y": line.bbox[1],
                    "line": line,
                })

        if urdu_container_group:
            min_y = min(l.bbox[1] for l in urdu_container_group)
            layout_items.append({
                "type": "urdu_container",
                "y": min_y,
                "lines": list(urdu_container_group),
            })
            urdu_container_group.clear()

        # 2. Add tables
        for tbl in tables:
            layout_items.append({
                "type": "table",
                "y": tbl.bbox[1],
                "table": tbl,
            })

        # 3. Add visual elements
        for vis in visuals:
            layout_items.append({
                "type": "visual",
                "y": vis.bbox[1],
                "visual": vis,
            })

        # 4. Partition items by reading block if document has repeated sections
        from pipeline.typography_engine import TypographyEngine
        split_y = TypographyEngine._detect_block_split(page)
        for item in layout_items:
            item["block_idx"] = 1 if (split_y is not None and item["y"] >= split_y) else 0

        # Sort items primarily by block_idx, then vertical coordinate
        layout_items.sort(key=lambda item: (item.get("block_idx", 0), item["y"]))

        # 5. Render sequentially with block boundary separation
        current_block = 0
        for item in layout_items:
            b_idx = item.get("block_idx", 0)
            if b_idx > current_block:
                current_block = b_idx
                # Insert subtle spacing paragraph between repeated voucher blocks
                p_sep = doc.add_paragraph()
                p_sep.paragraph_format.space_before = Pt(8)
                p_sep.paragraph_format.space_after = Pt(8)

            itype = item["type"]
            if itype == "line":
                self._render_text_line(doc, item["line"])
            elif itype == "urdu_container":
                self._render_urdu_container(doc, item["lines"])
            elif itype == "table":
                self._render_table(doc, item["table"], page_w)
            elif itype == "visual":
                self._render_visual(doc, item["visual"], page_w)

    def _render_text_line(self, doc: Document, line: FormattedLine) -> None:
        """Renders a single FormattedLine with typography, alignment, BiDi formatting, and underlines."""
        # Handle multi-column lines (e.g. Date & Month or 3-column signature labels)
        if line.multi_column_items and len(line.multi_column_items) > 1:
            self._render_multi_column_line(doc, line)
            return

        text = line.text.strip()
        if not text:
            return

        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15

        # Alignment
        if line.alignment == "center":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif line.alignment == "right":
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT

        font_size = max(6.0, min(72.0, float(line.font_size_pt or 10.0)))

        # Free-standing Urdu line (outside container)
        if line.language_script == "urdu_nastaliq" or ScriptClassifier.is_urdu_unicode(text):
            apply_bidi_paragraph_formatting(p)
            run = p.add_run(text)
            apply_bidi_run_formatting(
                run,
                font_cs=self.urdu_font_cs,
                font_ascii=self.urdu_font_ascii,
                font_size_pt=font_size,
                is_bold=line.is_bold,
            )
            return

        # Key-value preservation: render label and value on same unified paragraph
        if line.is_key_value and ":" in text:
            lbl, sep, val = text.partition(":")
            run_lbl = p.add_run(lbl + sep)
            run_lbl.font.name = self.default_font
            run_lbl.font.size = Pt(font_size)
            run_lbl.font.bold = True
            if line.is_underlined:
                run_lbl.font.underline = True
            if val:
                run_val = p.add_run(val)
                run_val.font.name = self.default_font
                run_val.font.size = Pt(font_size)
                run_val.font.bold = line.is_bold
                if line.is_underlined:
                    run_val.font.underline = True
        else:
            run = p.add_run(text)
            run.font.name = self.default_font
            run.font.size = Pt(font_size)
            run.font.bold = line.is_bold
            if line.is_underlined:
                run.font.underline = True

        # Form dotted fill line
        if line.has_fill_line:
            run_fill = p.add_run(" " + "." * 45)
            run_fill.font.name = self.default_font
            run_fill.font.size = Pt(font_size)
            run_fill.font.color.rgb = RGBColor(148, 163, 184)

    def _render_multi_column_line(self, doc: Document, line: FormattedLine) -> None:
        """Renders multi-column row (Date/Month or 3-column signatures) as a borderless layout table."""
        items = line.multi_column_items or []
        n_cols = len(items)
        if n_cols == 0:
            return

        tbl = doc.add_table(rows=1, cols=n_cols)
        set_table_no_borders(tbl)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False

        col_w = Inches(self.usable_width_in / n_cols)
        for ci, (col_text, col_align, is_u) in enumerate(items):
            tbl.columns[ci].width = col_w
            cell = tbl.cell(0, ci)
            cell.width = col_w
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            if col_align == "center":
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif col_align == "right":
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT

            run = p.add_run(col_text)
            run.font.name = self.default_font
            run.font.size = Pt(line.font_size_pt or 10.0)
            run.font.bold = line.is_bold
            if is_u or line.is_underlined:
                run.font.underline = True

    def _render_urdu_container(self, doc: Document, lines: List[FormattedLine]) -> None:
        """Renders boxed Urdu instructions as a 1x1 Table Grid callout container."""
        create_urdu_container_table(
            doc,
            lines,
            font_cs=self.urdu_font_cs,
            font_ascii=self.urdu_font_ascii,
        )
        p_sp = doc.add_paragraph()
        p_sp.paragraph_format.space_before = Pt(0)
        p_sp.paragraph_format.space_after = Pt(4)

    def _render_table(self, doc: Document, extracted_table: ExtractedTable, page_w: int) -> None:
        """Constructs a genuine Microsoft Word Table Grid matching extracted dimensions and cells."""
        n_rows = max(1, extracted_table.rows)
        n_cols = max(1, extracted_table.cols)

        w_table = doc.add_table(rows=n_rows, cols=n_cols)
        w_table.style = "Table Grid"
        w_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        w_table.autofit = False

        # 1. Proportional column widths
        if extracted_table.col_widths and len(extracted_table.col_widths) == n_cols:
            ratios = extracted_table.col_widths
            sum_r = sum(ratios)
            if sum_r > 0:
                ratios = [r / sum_r for r in ratios]
            else:
                ratios = [1.0 / n_cols] * n_cols
        else:
            ratios = [1.0 / n_cols] * n_cols

        for col_idx, ratio in enumerate(ratios):
            col_w = Inches(self.usable_width_in * ratio)
            w_table.columns[col_idx].width = col_w
            for row in w_table.rows:
                row.cells[col_idx].width = col_w

        # 2. Table borders & padding
        set_table_borders_and_margins(w_table, border_color=self.border_color)

        # 3. Merged cells resolution
        cell_map = {(c.row_idx, c.col_idx): c for c in extracted_table.cells}
        merged_cells = set()

        for (r, c), cell_data in cell_map.items():
            if cell_data.row_span > 1 or cell_data.col_span > 1:
                target_r = min(n_rows - 1, r + cell_data.row_span - 1)
                target_c = min(n_cols - 1, c + cell_data.col_span - 1)
                try:
                    top_left = w_table.cell(r, c)
                    bottom_right = w_table.cell(target_r, target_c)
                    top_left.merge(bottom_right)
                    for mr in range(r, target_r + 1):
                        for mc in range(c, target_c + 1):
                            if (mr, mc) != (r, c):
                                merged_cells.add((mr, mc))
                except Exception as exc:
                    logger.warning(f"Error merging cell ({r}, {c}) to ({target_r}, {target_c}): {exc}")

        # Ensure header row shading is set
        for col_idx in range(n_cols):
            if (0, col_idx) not in merged_cells:
                try:
                    set_cell_shading(w_table.cell(0, col_idx), self.header_shading)
                except Exception:
                    pass

        # 4. Cell content & formatting
        for (r, c), cell_data in cell_map.items():
            if (r, c) in merged_cells:
                continue
            try:
                cell = w_table.cell(r, c)
            except Exception:
                continue

            # Header row shading fill #EBF8FF
            if cell_data.is_header or r == 0:
                set_cell_shading(cell, self.header_shading)

            cell.text = ""
            lines = [ln.strip() for ln in cell_data.text.split("\n") if ln.strip()]
            if not lines:
                lines = [""]

            for li, ln in enumerate(lines):
                if li == 0:
                    p = cell.paragraphs[0]
                else:
                    p = cell.add_paragraph()

                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.line_spacing = 1.0

                # Alignment
                if cell_data.align == "center" or (cell_data.is_header or r == 0):
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                elif cell_data.align == "right":
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT

                # BiDi Urdu in cell
                if ScriptClassifier.is_urdu_unicode(ln):
                    apply_bidi_paragraph_formatting(p)
                    run = p.add_run(ln)
                    apply_bidi_run_formatting(
                        run,
                        font_cs=self.urdu_font_cs,
                        font_ascii=self.urdu_font_ascii,
                        font_size_pt=9.5 if cell_data.is_header else 9.0,
                        is_bold=cell_data.is_header or (r == 0),
                    )
                else:
                    run = p.add_run(ln)
                    run.font.name = self.default_font
                    run.font.size = Pt(9.5 if (cell_data.is_header or r == 0) else 9.0)
                    run.font.bold = cell_data.is_header or (r == 0)

        # Spacing paragraph after table
        p_sp = doc.add_paragraph()
        p_sp.paragraph_format.space_before = Pt(0)
        p_sp.paragraph_format.space_after = Pt(4)

    def _render_visual(self, doc: Document, visual: ExtractedVisualElement, page_w: int) -> None:
        """Embeds cropped RGB visual asset preserving aspect ratio and spatial positioning."""
        if visual.crop_image is None:
            return
        if visual.crop_image.width <= 0 or visual.crop_image.height <= 0:
            return

        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)

        # Horizontal alignment based on bounding box centroid
        cx = visual.bbox[0] + visual.bbox[2] / 2.0
        norm_cx = cx / max(1, page_w)
        if norm_cx < 0.35:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        elif norm_cx > 0.65:
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Width calculation preserving aspect ratio
        width_ratio = visual.bbox[2] / max(1, page_w)
        target_w_in = max(0.8, min(self.usable_width_in, self.usable_width_in * width_ratio * 1.2))

        # Size tuning per element type
        if visual.element_type == "photo":
            target_w_in = max(1.2, min(1.8, target_w_in))
        elif visual.element_type == "signature":
            target_w_in = max(1.2, min(2.2, target_w_in))
        elif visual.element_type == "barcode":
            target_w_in = max(1.5, min(2.5, target_w_in))
        elif visual.element_type == "logo":
            target_w_in = max(1.0, min(1.8, target_w_in))

        stream = pil_image_to_stream(visual.crop_image)
        run = p.add_run()
        run.add_picture(stream, width=Inches(target_w_in))
