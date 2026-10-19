"""
Shared Test Infrastructure & Fixtures for Compixor Image-to-Word E2E Test Suite.
Target Service URL: http://127.0.0.1:8003/api/convert-img-to-docx
"""

import io
import os
import sys
import time
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import Optional, Union

import docx
from docx.document import Document as DocxDocument
from docx.table import Table
from docx.text.paragraph import Paragraph
import numpy as np
from PIL import Image, ImageDraw, ImageFont
import pytest
import requests

# Namespaces used in OpenXML WordprocessingML
XML_NAMESPACES = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "pic": "http://schemas.openxmlformats.org/drawingml/2006/picture",
}

DEFAULT_BASE_URL = os.environ.get("IMG2WORD_TEST_API_URL", "http://127.0.0.1:8003").rstrip("/")
DEFAULT_CONVERT_ENDPOINT = f"{DEFAULT_BASE_URL}/api/convert-img-to-docx"
DEFAULT_HEALTH_ENDPOINT = f"{DEFAULT_BASE_URL}/health"

PRIMARY_CAPTURE_PNG = Path(r"C:\Users\AL GHANI\Desktop\Capture.PNG")
MIRROR_CAPTURE_PNG = Path(
    r"C:\Users\AL GHANI\.gemini\antigravity\brain\07c5f823-79a0-433f-8488-3bf2ec3bf7cb\.user_uploaded\media_1791036081237.png"
)


class DocxInspector:
    """Opaque-box validator for generated Microsoft Word (.docx) files."""

    def __init__(self, docx_bytes: bytes):
        self.raw_bytes = docx_bytes
        self.is_valid_zip = False
        self.namelist: list[str] = []
        self._doc_xml_str: Optional[str] = None
        self._doc_xml_tree: Optional[ET.Element] = None
        self._document: Optional[DocxDocument] = None

        self._inspect_zip()

    def _inspect_zip(self):
        try:
            with zipfile.ZipFile(io.BytesIO(self.raw_bytes)) as z:
                self.namelist = z.namelist()
                if "word/document.xml" in self.namelist:
                    self.is_valid_zip = True
                    self._doc_xml_str = z.read("word/document.xml").decode("utf-8", errors="replace")
                    self._doc_xml_tree = ET.fromstring(self._doc_xml_str)
        except Exception:
            self.is_valid_zip = False

        if self.is_valid_zip:
            try:
                self._document = docx.Document(io.BytesIO(self.raw_bytes))
            except Exception:
                self._document = None

    @property
    def document(self) -> DocxDocument:
        if self._document is None:
            raise ValueError("DOCX bytes could not be parsed as a valid python-docx Document")
        return self._document

    @property
    def raw_xml(self) -> str:
        return self._doc_xml_str or ""

    @property
    def xml_tree(self) -> Optional[ET.Element]:
        return self._doc_xml_tree

    def get_media_files(self) -> list[str]:
        """Returns all media files embedded in word/media/."""
        return [f for f in self.namelist if f.startswith("word/media/")]

    def get_media_count(self) -> int:
        return len(self.get_media_files())

    def get_media_bytes(self, filename: str) -> bytes:
        with zipfile.ZipFile(io.BytesIO(self.raw_bytes)) as z:
            return z.read(filename)

    def get_all_paragraphs_text(self) -> list[str]:
        """Extracts text from all Document paragraphs (top-level body paragraphs)."""
        if self._document is None:
            return []
        return [p.text for p in self._document.paragraphs]

    def get_body_text_full(self) -> str:
        """Returns concatenated text of all body paragraphs (excluding table cells)."""
        return "\n".join(self.get_all_paragraphs_text())

    def get_tables(self) -> list[Table]:
        if self._document is None:
            return []
        return list(self._document.tables)

    def get_table_data(self, table_index: int = 0) -> list[list[str]]:
        """Returns 2D grid matrix of text cells for the given table."""
        tables = self.get_tables()
        if table_index >= len(tables):
            raise IndexError(f"Table index {table_index} out of range (found {len(tables)} tables)")
        table = tables[table_index]
        matrix = []
        for row in table.rows:
            row_vals = [cell.text.strip() for cell in row.cells]
            matrix.append(row_vals)
        return matrix

    def has_bidi_tag(self) -> bool:
        """Checks for presence of w:bidi in OpenXML markup."""
        xml = self.raw_xml
        return ("<w:bidi" in xml) or ("w:bidi/>" in xml) or ("<w:rtl" in xml) or ("w:rtl/>" in xml)

    def has_font_binding(self, font_name: str) -> bool:
        """Checks whether the font name is referenced in document.xml font properties."""
        return font_name.lower() in self.raw_xml.lower()

    def find_paragraph(self, substr: str) -> Optional[Paragraph]:
        """Finds first paragraph containing substr (case-insensitive)."""
        substr_lower = substr.lower()
        for p in self._document.paragraphs:
            if substr_lower in p.text.lower():
                return p
        return None

    def find_paragraphs(self, substr: str) -> list[Paragraph]:
        """Finds all paragraphs containing substr (case-insensitive)."""
        substr_lower = substr.lower()
        return [p for p in self._document.paragraphs if substr_lower in p.text.lower()]

    def is_paragraph_centered(self, p: Paragraph) -> bool:
        """Inspects if paragraph has center alignment in python-docx or raw XML."""
        xml = p._p.xml
        return ('w:jc w:val="center"' in xml) or ('w:val="center"' in xml)

    def is_paragraph_right_aligned(self, p: Paragraph) -> bool:
        """Inspects if paragraph has right alignment in python-docx or raw XML."""
        xml = p._p.xml
        return ('w:jc w:val="right"' in xml) or ('w:val="right"' in xml)

    def is_paragraph_bold(self, p: Paragraph) -> bool:
        """Checks if any run in the paragraph or run properties has bold flag."""
        xml = p._p.xml
        return ("<w:b/>" in xml) or ('<w:b w:val="true"' in xml) or ('<w:b w:val="1"' in xml)

    def has_table_grid_style(self, table: Table) -> bool:
        """Checks if the table uses Table Grid or has explicit border markup."""
        xml = table._tbl.xml
        return ("TableGrid" in xml) or ("<w:tblBorders" in xml) or ("<w:tcBorders" in xml)

    def get_table_header_shading(self, table: Table) -> list[str]:
        """Extracts shading fill colors from the first row cells."""
        shd_colors = []
        if not table.rows:
            return shd_colors
        for cell in table.rows[0].cells:
            cell_xml = cell._tc.xml
            root = ET.fromstring(cell_xml)
            shd = root.find(".//w:shd", XML_NAMESPACES)
            if shd is not None and "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fill" in shd.attrib:
                shd_colors.append(shd.attrib["{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fill"])
            elif "w:fill" in cell_xml:
                # regex/heuristic fallback
                import re
                m = re.search(r'w:fill="([A-Fa-f0-9]+)"', cell_xml)
                if m:
                    shd_colors.append(m.group(1))
        return shd_colors


class ApiClient:
    """HTTP Client connecting to FastAPI conversion backend on port 8003."""

    def __init__(self, base_url: str = DEFAULT_BASE_URL, timeout: float = 60.0):
        self.base_url = base_url
        self.convert_url = f"{base_url}/api/convert-img-to-docx"
        self.health_url = f"{base_url}/health"
        self.timeout = timeout
        self.session = requests.Session()

    def is_alive(self) -> bool:
        try:
            r = self.session.get(self.health_url, timeout=3.0)
            return r.status_code == 200
        except Exception:
            return False

    def check_health(self) -> requests.Response:
        try:
            return self.session.get(self.health_url, timeout=self.timeout)
        except requests.exceptions.ConnectionError as exc:
            pytest.fail(
                f"Connection failed to {self.health_url}. Ensure server is running on port 8003: "
                f"uvicorn main:app --host 0.0.0.0 --port 8003. Error: {exc}"
            )

    def convert(
        self,
        file_bytes: bytes,
        filename: str = "document.png",
        mime_type: str = "image/png",
        language: str = "eng",
        extra_headers: Optional[dict] = None,
    ) -> requests.Response:
        """Sends multipart file conversion request."""
        files = {
            "files": (filename, file_bytes, mime_type),
        }
        data = {
            "language": language,
        }
        headers = extra_headers or {}
        try:
            start_t = time.perf_counter()
            resp = self.session.post(
                self.convert_url,
                files=files,
                data=data,
                headers=headers,
                timeout=self.timeout,
            )
            elapsed_ms = (time.perf_counter() - start_t) * 1000.0
            resp.client_elapsed_ms = elapsed_ms
            return resp
        except requests.exceptions.ConnectionError as exc:
            pytest.fail(
                f"Connection failed to {self.convert_url}. Ensure server is running on port 8003: "
                f"uvicorn main:app --host 0.0.0.0 --port 8003. Error: {exc}"
            )

    def convert_and_inspect(
        self,
        file_bytes: bytes,
        filename: str = "document.png",
        mime_type: str = "image/png",
        language: str = "eng",
        assert_status: int = 200,
    ) -> tuple[requests.Response, DocxInspector]:
        """Converts file and wraps output in DocxInspector."""
        resp = self.convert(file_bytes, filename=filename, mime_type=mime_type, language=language)
        assert resp.status_code == assert_status, (
            f"Expected status {assert_status}, got {resp.status_code}. Response: {resp.text[:500]}"
        )
        inspector = DocxInspector(resp.content)
        return resp, inspector


# ==============================================================================
# Pytest Fixtures
# ==============================================================================

@pytest.fixture(scope="session")
def api_client() -> ApiClient:
    """Fixture providing initialized API client targeting port 8003."""
    client = ApiClient()
    return client


@pytest.fixture(scope="session")
def admit_card_png_bytes() -> bytes:
    """Loads authoritative Capture.PNG test sample image."""
    for path in [PRIMARY_CAPTURE_PNG, MIRROR_CAPTURE_PNG]:
        if path.exists() and path.is_file():
            return path.read_bytes()

    # If neither exists, fail with actionable message
    pytest.fail(
        f"Test Sample 1 Capture.PNG not found at {PRIMARY_CAPTURE_PNG} or {MIRROR_CAPTURE_PNG}"
    )


@pytest.fixture(scope="session")
def sample_clean_text_image() -> bytes:
    """Creates a high-resolution clean synthetic text document image."""
    img = Image.new("RGB", (1200, 1600), color=(255, 255, 255))
    draw = ImageDraw.Draw(img)

    # Header
    draw.text((350, 100), "OFFICIAL CERTIFICATE", fill=(0, 0, 0))
    draw.text((450, 160), "DEPARTMENT OF EDUCATION", fill=(0, 0, 0))

    # Key-values
    draw.text((150, 300), "CANDIDATE NAME: AHMED BILAL", fill=(0, 0, 0))
    draw.text((150, 360), "FATHER NAME: MUHAMMAD BILAL", fill=(0, 0, 0))
    draw.text((150, 420), "REGISTRATION ID: 99482014", fill=(0, 0, 0))
    draw.text((150, 480), "PROGRAMME: COMPUTER SCIENCE", fill=(0, 0, 0))

    # Body
    draw.text(
        (150, 600),
        "This is to certify that the candidate has successfully completed the coursework.",
        fill=(0, 0, 0),
    )

    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


@pytest.fixture(scope="session")
def sample_table_image() -> bytes:
    """Creates a clean synthetic image containing a bordered 3-row x 3-column table."""
    img = Image.new("RGB", (1200, 1000), color=(255, 255, 255))
    draw = ImageDraw.Draw(img)

    draw.text((400, 50), "EXAMINATION RESULTS", fill=(0, 0, 0))

    # Table coordinates
    x0, y0 = 150, 150
    col_w = 300
    row_h = 80
    cols = 3
    rows = 3

    # Header fill
    draw.rectangle([x0, y0, x0 + cols * col_w, y0 + row_h], fill=(235, 248, 255))

    # Grid lines
    for r in range(rows + 1):
        y = y0 + r * row_h
        draw.line([(x0, y), (x0 + cols * col_w, y)], fill=(0, 0, 0), width=2)
    for c in range(cols + 1):
        x = x0 + c * col_w
        draw.line([(x, y0), (x, y0 + rows * row_h)], fill=(0, 0, 0), width=2)

    # Cell texts
    cells = [
        ["Course Code", "Course Title", "Grade"],
        ["CS-101", "Introduction to Programming", "A"],
        ["MATH-201", "Calculus and Analytic Geometry", "B+"],
    ]

    for r_idx, row in enumerate(cells):
        for c_idx, val in enumerate(row):
            draw.text((x0 + c_idx * col_w + 20, y0 + r_idx * row_h + 25), val, fill=(0, 0, 0))

    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


@pytest.fixture(scope="session")
def sample_marksheet_image() -> bytes:
    """Creates a comprehensive unseen synthetic academic marksheet document."""
    img = Image.new("RGB", (1400, 1800), color=(255, 255, 255))
    draw = ImageDraw.Draw(img)

    # Visual Logo element: drawn as a distinct circular badge in top-left
    draw.ellipse([80, 60, 200, 180], outline=(0, 51, 102), fill=(220, 235, 248), width=4)
    draw.text((115, 110), "BISE", fill=(0, 51, 102))

    # Centered Header
    draw.text((380, 80), "BOARD OF INTERMEDIATE & SECONDARY EDUCATION", fill=(0, 0, 0))
    draw.text((540, 130), "SECONDARY SCHOOL CERTIFICATE", fill=(0, 0, 0))
    draw.text((620, 175), "PROVISIONAL RESULT MARKSHEET", fill=(0, 0, 0))

    # Top right date
    draw.text((1150, 80), "Date: 15-Aug-2026", fill=(0, 0, 0))

    # Key-value metadata
    draw.text((100, 260), "ROLL NUMBER: 489201", fill=(0, 0, 0))
    draw.text((100, 310), "CANDIDATE NAME: MUHAMMAD USMAN", fill=(0, 0, 0))
    draw.text((100, 360), "FATHER NAME: TARIQ MAHMOOD", fill=(0, 0, 0))
    draw.text((100, 410), "INSTITUTION: GOVT HIGH SCHOOL LAHORE", fill=(0, 0, 0))

    # Student portrait photo placeholder (dark rectangular photo box on right)
    draw.rectangle([1120, 250, 1280, 430], outline=(100, 100, 100), fill=(180, 200, 220), width=2)
    draw.text((1165, 330), "PHOTO", fill=(50, 50, 50))

    # 5-Column Marksheet Table
    x0, y0 = 100, 500
    col_widths = [100, 380, 200, 220, 300]
    row_h = 60
    table_w = sum(col_widths)
    rows = 6

    # Header fill
    draw.rectangle([x0, y0, x0 + table_w, y0 + row_h], fill=(235, 248, 255))

    # Table Grid lines
    curr_y = y0
    for r in range(rows + 1):
        draw.line([(x0, curr_y), (x0 + table_w, curr_y)], fill=(0, 0, 0), width=2)
        curr_y += row_h

    curr_x = x0
    for w in col_widths:
        draw.line([(curr_x, y0), (curr_x, y0 + rows * row_h)], fill=(0, 0, 0), width=2)
        curr_x += w
    draw.line([(curr_x, y0), (curr_x, y0 + rows * row_h)], fill=(0, 0, 0), width=2)

    # 5-Column Table Content
    data = [
        ["S.No", "Subject", "Max Marks", "Marks Obtained", "Remarks"],
        ["1", "English Compulsory", "100", "88", "Pass"],
        ["2", "Urdu Compulsory", "100", "91", "Pass"],
        ["3", "Mathematics", "100", "96", "Pass"],
        ["4", "Physics Theory & Practical", "100", "85", "Pass"],
        ["5", "Chemistry Theory & Practical", "100", "89", "Pass"],
    ]

    for r_idx, row in enumerate(data):
        c_x = x0
        for c_idx, cell_val in enumerate(row):
            draw.text((c_x + 15, y0 + r_idx * row_h + 18), cell_val, fill=(0, 0, 0))
            c_x += col_widths[c_idx]

    # Summary
    draw.text((100, 920), "TOTAL MARKS: 449 / 500 (GRADE: A-ONE)", fill=(0, 0, 0))

    # Stamp / Signature visual flourish
    draw.ellipse([1050, 1000, 1220, 1120], outline=(150, 0, 0), width=3)
    draw.text((1090, 1050), "OFFICIAL SEAL", fill=(150, 0, 0))
    draw.line([(1050, 1180), (1250, 1180)], fill=(0, 0, 0), width=2)
    draw.text((1070, 1200), "CONTROLLER OF EXAMINATIONS", fill=(0, 0, 0))

    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


@pytest.fixture(scope="session")
def sample_cnic_image() -> bytes:
    """Creates a synthetic National Identity Card (CNIC) document image."""
    img = Image.new("RGB", (1000, 630), color=(240, 248, 245))
    draw = ImageDraw.Draw(img)

    # Green header band
    draw.rectangle([0, 0, 1000, 80], fill=(0, 100, 50))
    draw.text((250, 25), "GOVERNMENT OF PAKISTAN - NATIONAL IDENTITY CARD", fill=(255, 255, 255))

    # Citizen Portrait Photo
    draw.rectangle([60, 120, 260, 360], outline=(50, 50, 50), fill=(200, 220, 240), width=3)
    draw.text((120, 230), "CITIZEN PHOTO", fill=(0, 0, 0))

    # Text key-values
    draw.text((320, 130), "IDENTITY NUMBER: 35201-1234567-9", fill=(0, 0, 0))
    draw.text((320, 180), "NAME: HAMZA KHALID", fill=(0, 0, 0))
    draw.text((320, 230), "FATHER NAME: KHALID MEHMOOD", fill=(0, 0, 0))
    draw.text((320, 280), "GENDER: MALE", fill=(0, 0, 0))
    draw.text((320, 330), "COUNTRY OF STAY: PAKISTAN", fill=(0, 0, 0))
    draw.text((320, 380), "DATE OF BIRTH: 12-05-1998", fill=(0, 0, 0))
    draw.text((320, 430), "DATE OF ISSUE: 10-06-2020", fill=(0, 0, 0))
    draw.text((320, 480), "DATE OF EXPIRY: 10-06-2030", fill=(0, 0, 0))

    # Signature flourish
    draw.line([(80, 480), (240, 480)], fill=(0, 0, 0), width=2)
    draw.text((100, 495), "HOLDER SIGNATURE", fill=(0, 0, 0))

    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


@pytest.fixture(scope="session")
def sample_multipage_pdf_bytes(sample_clean_text_image: bytes, sample_table_image: bytes) -> bytes:
    """Generates a valid 2-page PDF document in memory using Pillow."""
    img1 = Image.open(io.BytesIO(sample_clean_text_image)).convert("RGB")
    img2 = Image.open(io.BytesIO(sample_table_image)).convert("RGB")

    buf = io.BytesIO()
    img1.save(buf, format="PDF", save_all=True, append_images=[img2])
    return buf.getvalue()
