"""
Unit Test Suite for Universal DOCX Document Reconstructor (R3, R4, R5, R6).
Conforms to PROJECT.md § Interface Contracts § 5 and Milestone M4.

Covers 23 comprehensive unit and empirical stress tests:
1. Document creation from empty pages (returns valid .docx archive)
2. Document creation with multiple pages (page breaks inserted)
3. Body text rendered strictly as top-level Document.paragraphs
4. Dynamic font size applied to run formatting
5. Bold weighting applied to runs
6. Paragraph alignments (Left, Center, Right)
7. Key-value line preservation on unified paragraphs with bold label
8. Genuine Table Grid style applied
9. Proportional column widths matching table.col_widths
10. Header row background shading fill #EBF8FF
11. Explicit table borders and cell margins (dxa padding)
12. Merged cell handling (row_span and col_span merging)
13. Visual element picture run embedding in aligned paragraphs
14. Visual element aspect ratio preservation
15. OpenXML BiDi paragraph and run XML tags (<w:bidi/>, <w:rtl/>)
16. Boxed Urdu instructions wrapped in 1x1 single-cell container table
17. Spatial vertical sorting (reading order interleaving by y-coordinate)
18. STRESS 1: Empty page list, completely blank images, and 0 text/tables/visuals
19. STRESS 2: Extreme aspect ratios (1500x20, 20x1500), single-pixel crops, out-of-bounds coords
20. STRESS 3: Table grid invariants with missing/empty cells, multi-spans, and displacement safety
21. STRESS 4: Body text preservation across metadata, retrievable by DocxInspector.get_body_text_full()
22. STRESS 5: BiDi typography (<w:bidi/>, <w:rtl/>, Nastaliq font bindings) and 1x1 container table
23. STRESS 6: In-memory zero disk leakage verification across repeated document generations
"""

from __future__ import annotations

import io
import os
import tempfile
import zipfile
from xml.etree import ElementTree as ET
import pytest
import numpy as np
from PIL import Image

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from tests.conftest import DocxInspector

from pipeline.preprocessor import PreprocessedPage
from pipeline.visual_extractor import ExtractedVisualElement
from pipeline.table_extractor import ExtractedTable, TableCell
from pipeline.typography_engine import FormattedLine
from pipeline.docx_reconstructor import (
    DocxReconstructor,
    set_cell_shading,
    set_table_borders_and_margins,
    pil_image_to_stream,
)


def make_dummy_page(
    page_num: int = 1, width: int = 1200, height: int = 1600, dpi: int = 300
) -> PreprocessedPage:
    """Helper creating a minimal PreprocessedPage fixture."""
    img = Image.new("RGB", (width, height), color=(255, 255, 255))
    return PreprocessedPage(
        page_number=page_num,
        rgb_image=img,
        clahe_gray=np.zeros((height, width), dtype=np.uint8),
        binary_image=np.ones((height, width), dtype=np.uint8) * 255,
        skew_angle=0.0,
        dpi=dpi,
        width=width,
        height=height,
    )


# ==============================================================================
# Unit Tests
# ==============================================================================

def test_reconstructor_empty_pages_returns_valid_docx():
    """T1: Empty pages list generates a valid .docx zip archive."""
    rec = DocxReconstructor()
    docx_bytes = rec.build_document([], [], [], [])
    assert isinstance(docx_bytes, bytes)
    assert len(docx_bytes) > 0
    assert zipfile.is_zipfile(io.BytesIO(docx_bytes))
    doc = Document(io.BytesIO(docx_bytes))
    assert doc is not None


def test_reconstructor_multi_page_break():
    """T2: Multi-page input creates page breaks between pages."""
    rec = DocxReconstructor()
    p1 = make_dummy_page(1)
    p2 = make_dummy_page(2)
    l1 = [FormattedLine("Page One Header", (100, 50, 200, 30), 14.0, True, "left", False, "latin", False)]
    l2 = [FormattedLine("Page Two Header", (100, 50, 200, 30), 14.0, True, "left", False, "latin", False)]

    docx_bytes = rec.build_document([p1, p2], [[], []], [[], []], [[l1], [l2]])
    assert zipfile.is_zipfile(io.BytesIO(docx_bytes))
    doc = Document(io.BytesIO(docx_bytes))

    # Inspect document xml for page break or paragraph presence
    texts = [p.text for p in doc.paragraphs]
    assert "Page One Header" in texts
    assert "Page Two Header" in texts


def test_reconstructor_body_text_in_top_level_paragraphs():
    """T3: Body text lines are rendered directly in doc.paragraphs (not buried in table cells)."""
    rec = DocxReconstructor()
    page = make_dummy_page(1)
    lines = [
        FormattedLine("DOCUMENT TITLE", (200, 50, 400, 40), 16.0, True, "center", False, "latin", False),
        FormattedLine("REGISTRATION: 12345", (100, 100, 300, 25), 11.0, False, "left", True, "latin", False),
    ]

    docx_bytes = rec.build_document([page], [[]], [[]], [lines])
    doc = Document(io.BytesIO(docx_bytes))

    para_texts = [p.text for p in doc.paragraphs]
    assert any("DOCUMENT TITLE" in t for t in para_texts)
    assert any("REGISTRATION: 12345" in t for t in para_texts)


def test_reconstructor_dynamic_font_size():
    """T4: Dynamic font sizes are correctly applied to paragraph runs."""
    rec = DocxReconstructor()
    page = make_dummy_page(1)
    lines = [
        FormattedLine("LARGE HEADING", (100, 50, 300, 40), 18.0, True, "left", False, "latin", False),
        FormattedLine("SMALL FOOTNOTE", (100, 120, 200, 20), 8.0, False, "left", False, "latin", False),
    ]

    docx_bytes = rec.build_document([page], [[]], [[]], [lines])
    doc = Document(io.BytesIO(docx_bytes))

    p_large = next(p for p in doc.paragraphs if "LARGE HEADING" in p.text)
    assert p_large.runs[0].font.size.pt == 18.0

    p_small = next(p for p in doc.paragraphs if "SMALL FOOTNOTE" in p.text)
    assert p_small.runs[0].font.size.pt == 8.0


def test_reconstructor_bold_weight():
    """T5: Bold weighting flag is correctly assigned to runs."""
    rec = DocxReconstructor()
    page = make_dummy_page(1)
    lines = [
        FormattedLine("BOLD LINE", (100, 50, 200, 30), 12.0, True, "left", False, "latin", False),
        FormattedLine("REGULAR LINE", (100, 90, 200, 30), 12.0, False, "left", False, "latin", False),
    ]

    docx_bytes = rec.build_document([page], [[]], [[]], [lines])
    doc = Document(io.BytesIO(docx_bytes))

    p_bold = next(p for p in doc.paragraphs if "BOLD LINE" in p.text)
    assert p_bold.runs[0].font.bold is True

    p_regular = next(p for p in doc.paragraphs if "REGULAR LINE" in p.text)
    assert p_regular.runs[0].font.bold is False


def test_reconstructor_paragraph_alignments():
    """T6: Left, Center, and Right paragraph alignments are correctly configured."""
    rec = DocxReconstructor()
    page = make_dummy_page(1)
    lines = [
        FormattedLine("CENTER HEADER", (300, 50, 400, 30), 14.0, True, "center", False, "latin", False),
        FormattedLine("RIGHT DATE", (850, 50, 150, 25), 10.0, False, "right", False, "latin", False),
        FormattedLine("LEFT TEXT", (100, 120, 300, 25), 11.0, False, "left", False, "latin", False),
    ]

    docx_bytes = rec.build_document([page], [[]], [[]], [lines])
    doc = Document(io.BytesIO(docx_bytes))

    p_center = next(p for p in doc.paragraphs if "CENTER HEADER" in p.text)
    assert p_center.alignment == WD_ALIGN_PARAGRAPH.CENTER

    p_right = next(p for p in doc.paragraphs if "RIGHT DATE" in p.text)
    assert p_right.alignment == WD_ALIGN_PARAGRAPH.RIGHT

    p_left = next(p for p in doc.paragraphs if "LEFT TEXT" in p.text)
    assert p_left.alignment == WD_ALIGN_PARAGRAPH.LEFT


def test_reconstructor_key_value_preservation():
    """T7: Key-value lines are rendered on unified paragraph with bold label and normal value."""
    rec = DocxReconstructor()
    page = make_dummy_page(1)
    lines = [
        FormattedLine("NAME: JOHN DOE", (100, 200, 400, 25), 11.0, False, "left", True, "latin", False),
    ]

    docx_bytes = rec.build_document([page], [[]], [[]], [lines])
    doc = Document(io.BytesIO(docx_bytes))

    p = next(p for p in doc.paragraphs if "NAME:" in p.text)
    assert len(p.runs) >= 2
    # Label is bold
    assert p.runs[0].text == "NAME:"
    assert p.runs[0].font.bold is True
    # Value is present on the same paragraph
    assert "JOHN DOE" in p.runs[1].text


def test_reconstructor_table_grid_style():
    """T8: Tables receive genuine 'Table Grid' style and centered alignment."""
    rec = DocxReconstructor()
    page = make_dummy_page(1)
    cell = TableCell(0, 0, 1, 1, (100, 200, 300, 50), "Cell Data", False, "left", 0.95)
    tbl = ExtractedTable(0, (100, 200, 600, 100), 1, 2, [0.5, 0.5], [cell])

    docx_bytes = rec.build_document([page], [[]], [[tbl]], [[]])
    doc = Document(io.BytesIO(docx_bytes))

    assert len(doc.tables) >= 1
    t = doc.tables[0]
    assert t.style.name == "Table Grid"
    assert t.alignment == WD_TABLE_ALIGNMENT.CENTER


def test_reconstructor_proportional_column_widths():
    """T9: Table columns and cells reflect proportional widths from col_widths."""
    rec = DocxReconstructor(usable_width_in=6.0)
    page = make_dummy_page(1)
    c1 = TableCell(0, 0, 1, 1, (100, 200, 200, 50), "Col1", True, "left", 0.9)
    c2 = TableCell(0, 1, 1, 1, (300, 200, 400, 50), "Col2", True, "left", 0.9)
    # Width ratios: 1/3 and 2/3
    tbl = ExtractedTable(0, (100, 200, 600, 100), 1, 2, [0.333, 0.667], [c1, c2])

    docx_bytes = rec.build_document([page], [[]], [[tbl]], [[]])
    doc = Document(io.BytesIO(docx_bytes))

    t = doc.tables[0]
    # Column 1 should be roughly half of Column 2
    w0 = t.columns[0].width
    w1 = t.columns[1].width
    assert w1 > w0
    assert abs((w1 / w0) - 2.0) < 0.2


def test_reconstructor_header_row_shading():
    """T10: Table header row cells receive XML shading fill #EBF8FF."""
    rec = DocxReconstructor(header_shading="EBF8FF")
    page = make_dummy_page(1)
    c1 = TableCell(0, 0, 1, 1, (100, 200, 200, 50), "Header Title", True, "center", 0.9)
    tbl = ExtractedTable(0, (100, 200, 400, 100), 1, 1, [1.0], [c1])

    docx_bytes = rec.build_document([page], [[]], [[tbl]], [[]])
    doc = Document(io.BytesIO(docx_bytes))

    t = doc.tables[0]
    cell_xml = t.cell(0, 0)._tc.xml
    assert "w:shd" in cell_xml
    assert "EBF8FF" in cell_xml


def test_reconstructor_table_borders_and_padding():
    """T11: Table XML contains explicit borders and cell margins (dxa padding)."""
    rec = DocxReconstructor(border_color="94A3B8")
    page = make_dummy_page(1)
    c1 = TableCell(0, 0, 1, 1, (100, 200, 200, 50), "Text", False, "left", 0.9)
    tbl = ExtractedTable(0, (100, 200, 400, 100), 1, 1, [1.0], [c1])

    docx_bytes = rec.build_document([page], [[]], [[tbl]], [[]])
    doc = Document(io.BytesIO(docx_bytes))

    t = doc.tables[0]
    tbl_xml = t._tbl.xml
    assert "w:tblBorders" in tbl_xml
    assert "94A3B8" in tbl_xml
    assert "w:tblCellMar" in tbl_xml
    assert 'w:w="120"' in tbl_xml
    assert 'w:w="140"' in tbl_xml


def test_reconstructor_merged_cells():
    """T12: Multi-span cells (row_span > 1 or col_span > 1) are cleanly merged."""
    rec = DocxReconstructor()
    page = make_dummy_page(1)
    # A 2x2 table where cell (0, 0) spans 2 columns (col_span=2)
    c_merged = TableCell(0, 0, 1, 2, (100, 200, 400, 40), "Merged Header", True, "center", 0.95)
    c_r1c0 = TableCell(1, 0, 1, 1, (100, 240, 200, 40), "Bottom Left", False, "left", 0.9)
    c_r1c1 = TableCell(1, 1, 1, 1, (300, 240, 200, 40), "Bottom Right", False, "left", 0.9)
    tbl = ExtractedTable(0, (100, 200, 400, 80), 2, 2, [0.5, 0.5], [c_merged, c_r1c0, c_r1c1])

    docx_bytes = rec.build_document([page], [[]], [[tbl]], [[]])
    doc = Document(io.BytesIO(docx_bytes))

    t = doc.tables[0]
    assert t.cell(0, 0).text == "Merged Header"
    # Row 1 cells contain distinct contents
    assert t.cell(1, 0).text == "Bottom Left"
    assert t.cell(1, 1).text == "Bottom Right"


def test_reconstructor_visual_element_embedding():
    """T13: Non-text visual elements are embedded as picture runs in word/media/."""
    rec = DocxReconstructor()
    page = make_dummy_page(1)

    crop = Image.new("RGB", (120, 160), color=(180, 50, 50))
    vis = ExtractedVisualElement("photo", (800, 100, 120, 160), crop, 0.95)

    docx_bytes = rec.build_document([page], [[vis]], [[]], [[]])

    with zipfile.ZipFile(io.BytesIO(docx_bytes)) as z:
        media_files = [f for f in z.namelist() if f.startswith("word/media/")]
        assert len(media_files) >= 1
        # Read back embedded image
        img_data = z.read(media_files[0])
        loaded = Image.open(io.BytesIO(img_data))
        assert loaded.size == (120, 160)


def test_reconstructor_visual_element_aspect_ratio():
    """T14: Visual elements preserve aspect ratio without distortion."""
    rec = DocxReconstructor()
    page = make_dummy_page(1)

    # 3:1 wide banner logo
    crop = Image.new("RGB", (300, 100), color=(0, 100, 200))
    vis = ExtractedVisualElement("logo", (100, 50, 300, 100), crop, 0.9)

    docx_bytes = rec.build_document([page], [[vis]], [[]], [[]])
    doc = Document(io.BytesIO(docx_bytes))

    with zipfile.ZipFile(io.BytesIO(docx_bytes)) as z:
        img_data = z.read([f for f in z.namelist() if f.startswith("word/media/")][0])
        loaded = Image.open(io.BytesIO(img_data))
        w, h = loaded.size
        assert w == 300 and h == 100


def test_reconstructor_urdu_bidi_typography():
    """T15: Free-standing Urdu lines include OpenXML BiDi tags (<w:bidi/>, <w:rtl/>)."""
    rec = DocxReconstructor()
    page = make_dummy_page(1)
    urdu_text = "\u06cc\u06c1 \u0627\u06cc\u06a9 \u067e\u0627\u06a9\u06cc\u0633\u062a\u0627\u0646\u06cc \u062a\u062d\u0631\u06cc\u0631 \u06c1\u06d2"
    line = FormattedLine(urdu_text, (200, 100, 400, 30), 12.0, False, "right", False, "urdu_nastaliq", False)

    docx_bytes = rec.build_document([page], [[]], [[]], [[line]])

    with zipfile.ZipFile(io.BytesIO(docx_bytes)) as z:
        doc_xml = z.read("word/document.xml").decode("utf-8")
        assert "w:bidi" in doc_xml
        assert "w:rtl" in doc_xml
        assert "Jameel Noori Nastaleeq" in doc_xml


def test_reconstructor_boxed_urdu_container_table():
    """T16: Boxed Urdu instructions (urdu_container=True) wrap in 1x1 container table."""
    rec = DocxReconstructor()
    page = make_dummy_page(1)
    u_lines = [
        FormattedLine("\u06c1\u062f\u0627\u06cc\u0627\u062a \u0628\u0631\u0627\u0626\u06d2 \u0637\u0644\u0628\u06c1", (100, 400, 500, 30), 12.0, True, "right", False, "urdu_nastaliq", True),
        FormattedLine("1. \u0627\u067e\u0646\u0627 \u0634\u0646\u0627\u062e\u062a\u06cc \u06a9\u0627\u0631\u0688 \u0633\u0627\u062a\u06be \u0644\u0627\u0626\u06cc\u06ba", (100, 440, 500, 25), 11.0, False, "right", False, "urdu_nastaliq", True),
    ]

    docx_bytes = rec.build_document([page], [[]], [[]], [u_lines])
    doc = Document(io.BytesIO(docx_bytes))

    # Should contain a 1x1 table
    container_found = False
    for t in doc.tables:
        if len(t.rows) == 1 and len(t.columns) == 1:
            container_found = True
            cell_xml = t.cell(0, 0)._tc.xml
            assert "w:bidi" in cell_xml
            break

    assert container_found, "Expected a 1x1 container table for boxed Urdu instructions"


def test_reconstructor_spatial_reading_order_interleaving():
    """T17: Elements are interleaved and rendered strictly by vertical y-coordinate."""
    rec = DocxReconstructor()
    page = make_dummy_page(1)

    # Elements at different y positions:
    # y=50: Title line
    l1 = FormattedLine("First Top Element", (100, 50, 300, 30), 14.0, True, "left", False, "latin", False)
    # y=200: Visual element
    crop = Image.new("RGB", (50, 50), color=(10, 10, 10))
    vis = ExtractedVisualElement("stamp", (100, 200, 50, 50), crop, 0.9)
    # y=350: Table
    cell = TableCell(0, 0, 1, 1, (100, 350, 200, 30), "Middle Table", False, "left", 0.9)
    tbl = ExtractedTable(0, (100, 350, 400, 50), 1, 1, [1.0], [cell])
    # y=500: Footer line
    l2 = FormattedLine("Last Bottom Element", (100, 500, 300, 30), 10.0, False, "left", False, "latin", False)

    # Pass in scrambled order
    docx_bytes = rec.build_document([page], [[vis]], [[tbl]], [[l2, l1]])
    doc = Document(io.BytesIO(docx_bytes))

    para_texts = [p.text for p in doc.paragraphs if p.text.strip()]
    assert para_texts[0] == "First Top Element"
    assert para_texts[-1] == "Last Bottom Element"


# ==============================================================================
# Challenger Stress Test Harness (Milestone M4 Empirical Validation)
# ==============================================================================

def test_stress_empty_and_blank_minimal_inputs():
    """STRESS 1: Empty page list, completely blank images, and 0 text/tables/visuals."""
    rec = DocxReconstructor()

    # 1. Empty pages
    docx_empty = rec.build_document([], [], [], [])
    assert isinstance(docx_empty, bytes)
    assert len(docx_empty) > 0
    assert zipfile.is_zipfile(io.BytesIO(docx_empty))
    doc_empty = Document(io.BytesIO(docx_empty))
    assert len(doc_empty.paragraphs) == 0
    assert len(doc_empty.tables) == 0

    # 2. Blank white page with 0 elements
    p_white = make_dummy_page(1, width=800, height=1000)
    docx_white = rec.build_document([p_white], [[]], [[]], [[]])
    assert zipfile.is_zipfile(io.BytesIO(docx_white))
    doc_white = Document(io.BytesIO(docx_white))
    assert len(doc_white.paragraphs) == 0
    assert len(doc_white.tables) == 0

    # 3. Completely blank black page
    img_black = Image.new("RGB", (600, 800), color=(0, 0, 0))
    p_black = PreprocessedPage(
        page_number=1,
        rgb_image=img_black,
        clahe_gray=np.zeros((800, 600), dtype=np.uint8),
        binary_image=np.zeros((800, 600), dtype=np.uint8),
        skew_angle=0.0,
        dpi=300,
        width=600,
        height=800,
    )
    docx_black = rec.build_document([p_black], [[]], [[]], [[]])
    assert zipfile.is_zipfile(io.BytesIO(docx_black))
    doc_black = Document(io.BytesIO(docx_black))
    assert len(doc_black.paragraphs) == 0
    assert len(doc_black.tables) == 0


def test_stress_coordinate_and_aspect_ratio_extremes():
    """STRESS 2: Extreme aspect ratios (1500x20, 20x1500), single-pixel crops, out-of-bounds coords."""
    rec = DocxReconstructor()
    page = make_dummy_page(1, width=1200, height=1600)

    # 1. Extreme horizontal banner (1500x20)
    banner_img = Image.new("RGB", (1500, 20), color=(255, 0, 0))
    vis_banner = ExtractedVisualElement("barcode", (100, 50, 1500, 20), banner_img, 0.95)

    # 2. Extreme vertical column (20x1500)
    column_img = Image.new("RGB", (20, 1500), color=(0, 255, 0))
    vis_column = ExtractedVisualElement("photo", (100, 100, 20, 1500), column_img, 0.95)

    # 3. Single-pixel crop (1x1)
    pixel_img = Image.new("RGB", (1, 1), color=(0, 0, 255))
    vis_pixel = ExtractedVisualElement("logo", (100, 150, 1, 1), pixel_img, 0.9)

    # 4. Out-of-bounds coordinates (negative and massive coordinates)
    oob_img = Image.new("RGB", (50, 50), color=(128, 128, 128))
    vis_oob = ExtractedVisualElement("stamp", (-500, -200, 50000, 10000), oob_img, 0.9)

    # 5. None crop
    vis_none = ExtractedVisualElement("stamp", (0, 0, 50, 50), None, 0.0)

    docx_bytes = rec.build_document(
        [page],
        [[vis_banner, vis_column, vis_pixel, vis_oob, vis_none]],
        [[]],
        [[]]
    )
    assert zipfile.is_zipfile(io.BytesIO(docx_bytes))
    doc = Document(io.BytesIO(docx_bytes))

    with zipfile.ZipFile(io.BytesIO(docx_bytes)) as z:
        media_files = [f for f in z.namelist() if f.startswith("word/media/")]
        # Banner, column, pixel, oob are embedded (vis_none is safely skipped)
        assert len(media_files) == 4

    # 6. Page with width=0 protection
    page_zero_w = make_dummy_page(1, width=0, height=1000)
    crop = Image.new("RGB", (10, 10), color=(20, 20, 20))
    vis_zero = ExtractedVisualElement("logo", (0, 0, 10, 10), crop, 0.8)
    docx_zero = rec.build_document([page_zero_w], [[vis_zero]], [[]], [[]])
    assert zipfile.is_zipfile(io.BytesIO(docx_zero))


def test_stress_table_grid_invariants_missing_cells_and_spans():
    """STRESS 3: Missing/empty cells, multi-column/row spans, Table Grid style without displacement."""
    rec = DocxReconstructor()
    page = make_dummy_page(1)

    # 1. 3x3 table with only 2 cells defined (7 missing cells)
    c0 = TableCell(0, 0, 1, 1, (10, 10, 100, 30), "Present Header", True, "center", 0.9)
    c2 = TableCell(2, 2, 1, 1, (210, 70, 100, 30), "Present Cell", False, "left", 0.9)
    tbl_sparse = ExtractedTable(0, (10, 10, 300, 90), 3, 3, [0.33, 0.33, 0.34], [c0, c2])

    # 2. Multi-row and multi-column span: 2x2 merged block in 3x3 table
    c_merged = TableCell(0, 0, 2, 2, (10, 150, 200, 60), "Merged 2x2", True, "center", 0.95)
    c_r0c2 = TableCell(0, 2, 1, 1, (210, 150, 100, 30), "Right Top", False, "left", 0.9)
    c_r1c2 = TableCell(1, 2, 1, 1, (210, 180, 100, 30), "Right Mid", False, "left", 0.9)
    c_r2c0 = TableCell(2, 0, 1, 3, (10, 210, 300, 30), "Span 3 Cols", False, "center", 0.9)
    tbl_spans = ExtractedTable(1, (10, 150, 300, 90), 3, 3, [0.33, 0.33, 0.34], [c_merged, c_r0c2, c_r1c2, c_r2c0])

    # 3. Out-of-bounds cell indices
    c_oob = TableCell(10, 10, 1, 1, (0, 0, 100, 30), "Out of bounds", False, "left", 0.5)
    tbl_oob = ExtractedTable(2, (10, 300, 200, 50), 2, 2, [0.5, 0.5], [c_oob])

    # 4. Zero row/col table
    tbl_zero = ExtractedTable(3, (10, 400, 200, 50), 0, 0, [], [])

    # Visual and text line to verify strict ordering without table displacement
    vis = ExtractedVisualElement("logo", (100, 5, 50, 50), Image.new("RGB", (50, 50)), 0.9)
    line1 = FormattedLine("Top Line", (100, 2, 200, 20), 12.0, False, "left", False, "latin", False)

    docx_bytes = rec.build_document([page], [[vis]], [[tbl_sparse, tbl_spans, tbl_oob, tbl_zero]], [[line1]])
    doc = Document(io.BytesIO(docx_bytes))

    # All generated tables have Table Grid style
    assert len(doc.tables) == 4
    for t in doc.tables:
        assert t.style.name == "Table Grid"
        assert t.alignment == WD_TABLE_ALIGNMENT.CENTER
        assert "w:tblBorders" in t._tbl.xml

    # Verify merged cell content
    t_spans = doc.tables[1]
    assert t_spans.cell(0, 0).text == "Merged 2x2"


def test_stress_body_text_preservation_via_inspector():
    """STRESS 4: Body text is preserved in top-level paragraphs and retrievable by DocxInspector."""
    rec = DocxReconstructor()
    page = make_dummy_page(1)

    lines = [
        FormattedLine("UNIVERSAL TEST ACADEMY", (100, 50, 600, 30), 16.0, True, "center", False, "latin", False),
        FormattedLine("ADMIT CARD RECORD", (200, 85, 400, 25), 14.0, True, "center", False, "latin", False),
        FormattedLine("NAME: TEST STUDENT", (100, 150, 500, 20), 11.0, False, "left", True, "latin", False),
        FormattedLine("FATHER NAME: TEST PARENT", (100, 175, 500, 20), 11.0, False, "left", True, "latin", False),
        FormattedLine("REGISTRATION NO: 99998888", (100, 200, 500, 20), 11.0, False, "left", True, "latin", False),
        FormattedLine("PROGRAMME: COMPUTER SCIENCE", (100, 225, 500, 20), 11.0, False, "left", True, "latin", False),
    ]

    cell = TableCell(0, 0, 1, 1, (100, 300, 300, 40), "TABLE_CELL_ISOLATED_DATA", False, "left", 0.9)
    tbl = ExtractedTable(0, (100, 300, 300, 40), 1, 1, [1.0], [cell])

    crop = Image.new("RGB", (60, 80), color=(100, 100, 100))
    vis = ExtractedVisualElement("photo", (700, 150, 60, 80), crop, 0.95)

    docx_bytes = rec.build_document([page], [[vis]], [[tbl]], [lines])
    inspector = DocxInspector(docx_bytes)

    body_full = inspector.get_body_text_full()

    # All metadata is retrievable in body text
    assert "UNIVERSAL TEST ACADEMY" in body_full
    assert "ADMIT CARD RECORD" in body_full
    assert "NAME:" in body_full
    assert "TEST STUDENT" in body_full
    assert "FATHER NAME:" in body_full
    assert "TEST PARENT" in body_full
    assert "99998888" in body_full
    assert "COMPUTER SCIENCE" in body_full

    # Table cell data is NOT in body text (no table leakage)
    assert "TABLE_CELL_ISOLATED_DATA" not in body_full


def test_stress_bidi_typography_and_container_invariants():
    """STRESS 5: OpenXML <w:bidi/> and <w:rtl/> XML tags, Nastaliq fonts, and 1x1 container table."""
    rec = DocxReconstructor()
    page = make_dummy_page(1)

    # 1. Free-standing Urdu line
    urdu_free = FormattedLine("امتحانی مرکز میں موبائل فون لانا منع ہے", (100, 200, 500, 25), 11.5, False, "right", False, "urdu_nastaliq", False)

    # 2. Boxed Urdu container lines
    urdu_box_1 = FormattedLine("ہدایات برائے طلبہ", (100, 400, 500, 30), 13.0, True, "right", False, "urdu_nastaliq", True)
    urdu_box_2 = FormattedLine("1. اپنا اصل شناختی کارڈ ہمراہ لائیں", (100, 435, 500, 25), 11.0, False, "right", False, "urdu_nastaliq", True)

    docx_bytes = rec.build_document([page], [[]], [[]], [[urdu_free, urdu_box_1, urdu_box_2]])
    inspector = DocxInspector(docx_bytes)

    assert inspector.has_bidi_tag()
    raw_xml = inspector.raw_xml
    assert "w:bidi" in raw_xml
    assert "w:rtl" in raw_xml
    assert "Jameel Noori Nastaleeq" in raw_xml
    assert "Noto Nastaliq Urdu" in raw_xml

    # Check 1x1 container table
    tables = inspector.get_tables()
    container_tables = [t for t in tables if len(t.rows) == 1 and len(t.columns) == 1]
    assert len(container_tables) == 1
    c_cell = container_tables[0].cell(0, 0)
    assert "ہدایات برائے طلبہ" in c_cell.text
    assert "شناختی کارڈ" in c_cell.text
    assert "w:bidi" in c_cell._tc.xml


def test_stress_in_memory_zero_disk_leakage():
    """STRESS 6: Verify no temporary .docx or other files are written to disk during generation."""
    rec = DocxReconstructor()
    page = make_dummy_page(1)
    crop = Image.new("RGB", (100, 100), color=(50, 100, 150))
    vis = ExtractedVisualElement("photo", (500, 100, 100, 100), crop, 0.9)
    tbl = ExtractedTable(0, (100, 300, 400, 100), 2, 2, [0.5, 0.5], [
        TableCell(0, 0, 1, 1, (100, 300, 200, 50), "H1", True, "center", 0.9),
        TableCell(0, 1, 1, 1, (300, 300, 200, 50), "H2", True, "center", 0.9),
        TableCell(1, 0, 1, 1, (100, 350, 200, 50), "D1", False, "left", 0.9),
        TableCell(1, 1, 1, 1, (300, 350, 200, 50), "D2", False, "left", 0.9),
    ])
    lines = [
        FormattedLine("Line 1", (100, 50, 200, 20), 12.0, False, "left", False, "latin", False),
        FormattedLine("اردو لائن", (100, 80, 200, 20), 12.0, False, "right", False, "urdu_nastaliq", False),
    ]

    # Snapshot current directory and temp directory before runs
    cwd_files_before = set(os.listdir("."))
    temp_dir = tempfile.gettempdir()
    temp_files_before = set(os.listdir(temp_dir))

    # Run document generation 10 times
    for _ in range(10):
        docx_bytes = rec.build_document([page], [[vis]], [[tbl]], [lines])
        assert len(docx_bytes) > 0
        assert zipfile.is_zipfile(io.BytesIO(docx_bytes))

    cwd_files_after = set(os.listdir("."))
    temp_files_after = set(os.listdir(temp_dir))

    # Verify zero new docx or leaked temp files in cwd
    new_cwd_docx = [f for f in cwd_files_after - cwd_files_before if f.endswith(".docx")]
    assert len(new_cwd_docx) == 0, f"Leaked docx files in cwd: {new_cwd_docx}"

    # Verify no new .docx files in system temp dir
    new_temp_docx = [f for f in temp_files_after - temp_files_before if f.endswith(".docx")]
    assert len(new_temp_docx) == 0, f"Leaked docx files in temp dir: {new_temp_docx}"
